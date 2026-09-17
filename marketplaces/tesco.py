
import json
import os
import re
import html
import requests

from openpyxl import load_workbook

from models import MarketplaceOrder
from marketplaces.base import MarketplaceAdapter


from config import (
    TESCO_MIRAKL_URL,
    TESCO_MIRAKL_API_KEY,
)


# ============================================================
# TESCO EXCEL IMPORT
# ============================================================

def build_tesco_rp2_search_ref(order_id):
    """
    Convert the full Tesco/Mirakl order number into the
    12-character reference used to locate the order in RPii.

    Examples:

        4431-2696-2301-A
        -> 4431-2696-23

        5431-2680-54-A
        -> 5431-2680-54

    IMPORTANT:
    The full Tesco order number remains the authoritative
    marketplace order ID. This shortened value is used only
    when searching RPii.
    """

    order_id = str(order_id).strip()

    if len(order_id) < 12:
        raise ValueError(
            f"Tesco order ID is too short to create "
            f"an RPii search reference: {order_id}"
        )

    return order_id[:12]

def read_tesco_orders(filename):
    """
    Read all orders from the Tesco Mirakl Excel export.

    Returns a list of MarketplaceOrder objects
    in spreadsheet order.

    This function is READ ONLY.
    """

    workbook = load_workbook(
        filename,
        data_only=True,
    )

    sheet = workbook.active

    headers = {
        cell.value: cell.column
        for cell in sheet[1]
        if cell.value
    }

    required_columns = [
        "Order number",
        "Offer SKU",
        "Amount",
        "Shipping address zip",
        "Shipping address phone",
        "Shipping address phone 2",
    ]

    for column in required_columns:

        if column not in headers:

            raise RuntimeError(
                f"Missing required Tesco Excel column: "
                f"{column}"
            )

    orders = []

    for row in range(
        2,
        sheet.max_row + 1,
    ):

        order_id = sheet.cell(
            row=row,
            column=headers["Order number"],
        ).value

        # Ignore completely empty rows.
        if not order_id:
            continue

        sku = sheet.cell(
            row=row,
            column=headers["Offer SKU"],
        ).value

        price = sheet.cell(
            row=row,
            column=headers["Amount"],
        ).value

        postcode = sheet.cell(
            row=row,
            column=headers["Shipping address zip"],
        ).value

        phone_1 = sheet.cell(
            row=row,
            column=headers["Shipping address phone"],
        ).value

        phone_2 = sheet.cell(
            row=row,
            column=headers["Shipping address phone 2"],
        ).value

        order = MarketplaceOrder(
            marketplace="TESCO",

            order_id=str(order_id).strip(),

            sku=str(sku).strip(),

            price=float(price),

            postcode=str(postcode).strip(),

            phone_1=(
                str(phone_1).strip()
                if phone_1
                else ""
            ),

            phone_2=(
                str(phone_2).strip()
                if phone_2
                else ""
            ),
        )

        orders.append(order)

    if not orders:
        raise RuntimeError(
            "No Tesco orders were found in the spreadsheet."
        )

    return orders


def read_first_tesco_order(filename):
    """
    Compatibility wrapper.

    Return the first order from the Tesco spreadsheet.
    """

    orders = read_tesco_orders(filename)

    return orders[0]



# ============================================================
# TESCO MIRAKL LIVE ORDER QUEUE
# ============================================================

def read_tesco_shipping_orders():
    """
    Read all Tesco Mirakl orders currently in SHIPPING state.

    READ ONLY.

    Safety behaviour:
    - Uses GET requests only.
    - Handles Mirakl pagination.
    - Returns the same MarketplaceOrder objects used by the
      existing Excel import route.
    - Fails closed if an order cannot be represented safely.
    """

    if not TESCO_MIRAKL_API_KEY:
        raise RuntimeError("TESCO_MIRAKL_API_KEY is missing.")

    url = f"{TESCO_MIRAKL_URL.rstrip('/')}/api/orders"

    headers = {
        "Authorization": TESCO_MIRAKL_API_KEY,
        "Accept": "application/json",
    }

    max_results = 100
    offset = 0
    marketplace_orders = []
    seen_order_ids = set()

    print()
    print("=" * 70)
    print("TESCO MIRAKL ORDER FETCH - READ ONLY")
    print("=" * 70)
    print("Fetching orders currently in SHIPPING state...")

    while True:
        params = {
            "order_state_codes": "SHIPPING",
            "max": max_results,
            "offset": offset,
        }

        response = requests.get(
            url,
            headers=headers,
            params=params,
            timeout=30,
        )

        if response.status_code != 200:
            raise RuntimeError(
                "Tesco Mirakl SHIPPING order fetch failed. "
                f"HTTP {response.status_code}: {response.text}"
            )

        try:
            data = response.json()
        except ValueError as exc:
            raise RuntimeError(
                "Tesco Mirakl SHIPPING order fetch did not return valid JSON."
            ) from exc

        raw_orders = data.get("orders")

        if not isinstance(raw_orders, list):
            raise RuntimeError(
                "Tesco Mirakl SHIPPING order fetch returned an unexpected "
                "response: 'orders' is not a list."
            )

        for raw_order in raw_orders:
            if not isinstance(raw_order, dict):
                raise RuntimeError(
                    "Tesco Mirakl SHIPPING order fetch returned an "
                    "unexpected order structure."
                )

            order_id = str(raw_order.get("order_id") or "").strip()
            order_state = str(raw_order.get("order_state") or "").strip().upper()

            if not order_id:
                raise RuntimeError(
                    "Tesco Mirakl returned a SHIPPING order without an order_id."
                )

            if order_state != "SHIPPING":
                raise RuntimeError(
                    f"Tesco Mirakl returned order {order_id} with unexpected "
                    f"state {order_state!r}. Expected SHIPPING."
                )

            if order_id in seen_order_ids:
                raise RuntimeError(
                    f"Tesco Mirakl returned duplicate order {order_id} "
                    "while paging the queue."
                )

            order_lines = raw_order.get("order_lines", [])
            if not isinstance(order_lines, list) or not order_lines:
                raise RuntimeError(
                    f"Tesco Mirakl order {order_id} contains no order lines."
                )

            # The current Tesco processor validates/processes one product
            # line. Do not silently choose a line if Mirakl returns more.
            if len(order_lines) != 1:
                raise RuntimeError(
                    f"Tesco Mirakl order {order_id} contains "
                    f"{len(order_lines)} order lines. "
                    "The current Tesco processor expects exactly 1."
                )

            line = order_lines[0]
            sku = str(line.get("offer_sku") or "").strip()
            if not sku:
                raise RuntimeError(
                    f"Tesco Mirakl order {order_id} has no offer SKU."
                )

            customer = raw_order.get("customer") or {}
            if not isinstance(customer, dict):
                raise RuntimeError(
                    f"Tesco Mirakl order {order_id} has invalid customer data."
                )

            shipping = customer.get("shipping_address") or {}
            if not isinstance(shipping, dict):
                raise RuntimeError(
                    f"Tesco Mirakl order {order_id} has invalid shipping data."
                )

            postcode = str(shipping.get("zip_code") or "").strip()
            if not postcode:
                raise RuntimeError(
                    f"Tesco Mirakl order {order_id} has no shipping postcode."
                )

            total_price = raw_order.get("total_price")
            try:
                price = float(total_price)
            except (TypeError, ValueError) as exc:
                raise RuntimeError(
                    f"Tesco Mirakl order {order_id} has an invalid total "
                    f"price: {total_price!r}."
                ) from exc

            marketplace_orders.append(
                MarketplaceOrder(
                    marketplace="TESCO",
                    order_id=order_id,
                    sku=sku,
                    price=price,
                    postcode=postcode,
                    phone_1=str(shipping.get("phone") or "").strip(),
                    phone_2=str(shipping.get("phone_secondary") or "").strip(),
                )
            )
            seen_order_ids.add(order_id)

        print(f"  Page offset {offset}: {len(raw_orders)} order(s)")

        if len(raw_orders) < max_results:
            break

        offset += max_results

    print()
    print(
        f"✓ Tesco Mirakl returned {len(marketplace_orders)} "
        "order(s) awaiting shipment."
    )
    print("No Tesco Mirakl write has been made.")
    print("=" * 70)

    return marketplace_orders


# ============================================================
# TESCO MARKETPLACE ADAPTER
# ============================================================

class TescoMarketplaceAdapter(MarketplaceAdapter):
    """
    Tesco marketplace adapter.

    Initial implementation reads orders from the Tesco
    Mirakl Excel export only.

    No Mirakl API writes, RPii actions, SMS actions or
    Stream actions are performed here.
    """

    def __init__(self, filename):
        self.filename = filename

    @property
    def name(self) -> str:
        return "Tesco"

    def read_first_order(self) -> MarketplaceOrder:
        return read_first_tesco_order(
            self.filename
        )

def read_tesco_mirakl_order(order_id):

    if not TESCO_MIRAKL_API_KEY:
        raise RuntimeError(
            "TESCO_MIRAKL_API_KEY is missing."
        )

    url = (
        f"{TESCO_MIRAKL_URL}"
        f"/api/orders"
    )

    headers = {
        "Authorization": TESCO_MIRAKL_API_KEY,
        "Accept": "application/json",
    }

    params = {
        "order_ids": order_id,
    }

    response = requests.get(
        url,
        headers=headers,
        params=params,
        timeout=30,
    )

    response.raise_for_status()

    data = response.json()

    orders = data.get(
        "orders",
        []
    )

    if not orders:
        return None

    return orders[0]

# ============================================================
# TESCO MIRAKL CUSTOMER COMMUNICATION - READ ONLY / DRY RUN
# ============================================================

TESCO_DELIVERY_TOPIC_TYPE = "FREE_TEXT"
TESCO_DELIVERY_SUBJECT = (
    "Electrical Discount UK - "
    "Delivery Information & Copy Invoice"
)


def clean_tesco_mirakl_message_body(body):
    """Convert Mirakl HTML message content into readable terminal text."""

    if not body:
        return ""

    text = str(body)

    text = re.sub(
        r"<br\s*/?>",
        "\n",
        text,
        flags=re.IGNORECASE,
    )

    text = re.sub(
        r"<[^>]+>",
        "",
        text,
    )

    text = html.unescape(text)

    return text.strip()


def read_tesco_mirakl_threads(order_id):
    """
    Read existing Tesco Mirakl threads/messages for an order.

    READ ONLY. This function performs GET requests only.
    """

    if not TESCO_MIRAKL_API_KEY:
        raise RuntimeError(
            "TESCO_MIRAKL_API_KEY is missing."
        )

    url = (
        f"{TESCO_MIRAKL_URL.rstrip('/')}"
        f"/api/inbox/threads"
    )

    headers = {
        "Authorization": TESCO_MIRAKL_API_KEY,
        "Accept": "application/json",
    }

    params = {
        "entity_type": "MMP_ORDER",
        "entity_id": order_id,
        "with_messages": "true",
    }

    response = requests.get(
        url,
        headers=headers,
        params=params,
        timeout=30,
    )

    if response.status_code != 200:
        raise RuntimeError(
            f"Tesco Mirakl thread lookup failed. "
            f"HTTP {response.status_code}: "
            f"{response.text}"
        )

    data = response.json()
    threads = data.get("data", [])
    all_messages = []

    for thread in threads:
        topic = thread.get("topic", {})

        for message in thread.get("messages", []):
            sender = message.get("from", {})

            all_messages.append(
                {
                    "thread_id": thread.get("id"),
                    "topic_type": topic.get("type"),
                    "topic_value": topic.get("value"),
                    # Kept as an alias for older test code.
                    "topic_code": topic.get("value"),
                    "message_id": message.get("id"),
                    "date": message.get("date_created"),
                    "sender_type": sender.get("type"),
                    "sender_name": sender.get("display_name"),
                    "attachments": message.get("attachments", []),
                    "body": clean_tesco_mirakl_message_body(
                        message.get("body")
                    ),
                }
            )

    all_messages.sort(
        key=lambda message: message["date"] or ""
    )

    return {
        "threads": threads,
        "messages": all_messages,
    }


def build_tesco_mirakl_dry_run_payload(
    order_id,
    customer_name,
    message_body,
    invoice_path,
    existing_threads,
):
    """
    Build a LOCAL representation of the proposed Tesco Mirakl
    delivery message. No network write occurs here.
    """

    thread_id = None

    if existing_threads:
        thread_id = existing_threads[0].get("id")

    return {
        "mode": "DRY_RUN",
        "order_id": order_id,
        "customer": customer_name,
        "topic": {
            "type": TESCO_DELIVERY_TOPIC_TYPE,
            "value": TESCO_DELIVERY_SUBJECT,
            "label": TESCO_DELIVERY_SUBJECT,
        },
        "existing_thread_id": thread_id,
        "message_body": message_body,
        "attachment": {
            "filename": os.path.basename(invoice_path),
            "path": invoice_path,
            "exists": os.path.isfile(invoice_path),
        },
    }


def send_tesco_mirakl_delivery_message(
    order_id,
    message_body,
    invoice_path,
):
    """
    Create ONE Tesco Mirakl order thread containing the delivery
    message and invoice PDF.

    IMPORTANT SAFETY DESIGN:
    - This function performs exactly one POST request.
    - It contains no retry loop.
    - The caller must perform all validation, duplicate checks and
      the exact SEND <order_id> operator confirmation first.
    """

    if not TESCO_MIRAKL_API_KEY:
        raise RuntimeError(
            "TESCO_MIRAKL_API_KEY is missing."
        )

    order_id = str(order_id).strip()

    if not order_id:
        raise RuntimeError(
            "Tesco order ID is blank."
        )

    if not str(message_body).strip():
        raise RuntimeError(
            "Tesco Mirakl message body is blank."
        )

    if not os.path.isfile(invoice_path):
        raise RuntimeError(
            f"Invoice attachment does not exist: {invoice_path}"
        )

    url = (
        f"{TESCO_MIRAKL_URL.rstrip('/')}"
        f"/api/orders/{order_id}/threads"
    )

    headers = {
        "Authorization": TESCO_MIRAKL_API_KEY,
        "Accept": "application/json",
    }

    # Tesco Mirakl's customer-facing UI collapses plain newline
    # characters when rendering a message. Convert them to HTML
    # line breaks for the POST payload only; the original message
    # remains unchanged for previews and safety verification.
    mirakl_message_body = (
        str(message_body)
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .replace("\n", "<br>")
    )

    thread_input = {
        "body": mirakl_message_body,
        "topic": {
            "type": TESCO_DELIVERY_TOPIC_TYPE,
            "value": TESCO_DELIVERY_SUBJECT,
        },
        # The Tesco order thread is intended for the customer.
        "to": ["CUSTOMER"],
    }

    with open(invoice_path, "rb") as invoice_file:
        files = [
            (
                "files",
                (
                    os.path.basename(invoice_path),
                    invoice_file,
                    "application/pdf",
                ),
            ),
            (
                "thread_input",
                (
                    None,
                    json.dumps(thread_input),
                    "application/json",
                ),
            ),
        ]

        # EXACTLY ONE WRITE ATTEMPT. DO NOT ADD RETRIES HERE.
        response = requests.post(
            url,
            headers=headers,
            files=files,
            timeout=30,
        )

    response_text = response.text

    if not (200 <= response.status_code < 300):
        raise RuntimeError(
            "Tesco Mirakl message POST failed. "
            f"HTTP {response.status_code}: {response_text}"
        )

    try:
        response_json = response.json()
    except ValueError:
        response_json = None

    return {
        "status_code": response.status_code,
        "response_json": response_json,
        "response_text": response_text,
        "endpoint": url,
        "topic_type": TESCO_DELIVERY_TOPIC_TYPE,
        "topic_value": TESCO_DELIVERY_SUBJECT,
        "attachment_filename": os.path.basename(invoice_path),
    }
