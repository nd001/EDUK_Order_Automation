import html
import os
import re
import json

import requests
from openpyxl import load_workbook

from config import (
    MIRAKL_BASE_URL,
    MIRAKL_API_KEY,
    DIY_DELIVERY_TOPIC_CODE,
)

from models import MarketplaceOrder
from marketplaces.base import MarketplaceAdapter

DIY_DIRECT_DELIVERY_CARRIER_CODE = "DIR"

DIY_DIRECT_DELIVERY_CARRIER_LABEL = (
    "The verified seller will contact you directly "
    "to plan a delivery"
)

# ============================================================
# DIY / B&Q EXCEL IMPORT
# ============================================================

def read_bq_orders(filename):
    """
    Read all orders from the B&Q unshipped-orders
    Excel export.

    Returns a list of MarketplaceOrder objects
    in spreadsheet order.
    """

    workbook = load_workbook(
        filename,
        data_only=True
    )

    sheet = workbook.active

    headers = {
        cell.value: cell.column
        for cell in sheet[1]
        if cell.value
    }

    required_columns = [
        "Order number",
        "Offer SKU",
        "Amount",
        "Shipping address zip",
        "Shipping address phone",
        "Shipping address phone 2",
    ]

    for column in required_columns:

        if column not in headers:

            raise RuntimeError(
                f"Missing required B&Q Excel column: "
                f"{column}"
            )

    orders = []

    for row in range(
        2,
        sheet.max_row + 1
    ):

        order_id = sheet.cell(
            row=row,
            column=headers["Order number"]
        ).value

        # Ignore completely empty rows.
        if not order_id:
            continue

        sku = sheet.cell(
            row=row,
            column=headers["Offer SKU"]
        ).value

        price = sheet.cell(
            row=row,
            column=headers["Amount"]
        ).value

        postcode = sheet.cell(
            row=row,
            column=headers["Shipping address zip"]
        ).value

        phone_1 = sheet.cell(
            row=row,
            column=headers["Shipping address phone"]
        ).value

        phone_2 = sheet.cell(
            row=row,
            column=headers["Shipping address phone 2"]
        ).value

        order = MarketplaceOrder(
            marketplace="DIY",

            order_id=str(order_id).strip(),

            sku=str(sku).strip(),

            price=float(price),

            postcode=str(postcode).strip(),

            phone_1=(
                str(phone_1).strip()
                if phone_1
                else ""
            ),

            phone_2=(
                str(phone_2).strip()
                if phone_2
                else ""
            ),
        )

        orders.append(order)

    if not orders:
        raise RuntimeError(
            "No B&Q orders were found in the spreadsheet."
        )

    return orders


def read_first_bq_order(filename):
    """
    Compatibility wrapper.

    Return the first order from the B&Q spreadsheet.
    """

    orders = read_bq_orders(filename)

    return orders[0]    


# ============================================================
# MIRAKL SHIPPING QUEUE
# ============================================================

def read_mirakl_shipping_orders():
    """
    Read all B&Q / DIY Mirakl orders currently in SHIPPING state.

    READ ONLY.

    Safety behaviour:
    - Uses GET requests only.
    - Handles Mirakl pagination.
    - Returns the same MarketplaceOrder objects used by the
      existing Excel import route.
    - Fails closed if an order cannot be represented safely.
    """

    url = (
        f"{MIRAKL_BASE_URL.rstrip('/')}"
        f"/api/orders"
    )

    headers = {
        "Authorization": MIRAKL_API_KEY,
        "Accept": "application/json",
    }

    max_results = 100
    offset = 0
    marketplace_orders = []
    seen_order_ids = set()

    print()
    print("=" * 70)
    print("B&Q MIRAKL ORDER FETCH - READ ONLY")
    print("=" * 70)
    print("Fetching orders currently in SHIPPING state...")

    while True:
        params = {
            "order_state_codes": "SHIPPING",
            "max": max_results,
            "offset": offset,
        }

        response = requests.get(
            url,
            headers=headers,
            params=params,
            timeout=30,
        )

        if response.status_code != 200:
            raise RuntimeError(
                "Mirakl SHIPPING order fetch failed. "
                f"HTTP {response.status_code}: "
                f"{response.text}"
            )

        try:
            data = response.json()
        except ValueError as exc:
            raise RuntimeError(
                "Mirakl SHIPPING order fetch did not return "
                "valid JSON."
            ) from exc

        raw_orders = data.get("orders")

        if not isinstance(raw_orders, list):
            raise RuntimeError(
                "Mirakl SHIPPING order fetch returned an "
                "unexpected response: 'orders' is not a list."
            )

        for raw_order in raw_orders:
            order_id = str(
                raw_order.get("order_id") or ""
            ).strip()

            order_state = str(
                raw_order.get("order_state") or ""
            ).strip().upper()

            if not order_id:
                raise RuntimeError(
                    "Mirakl returned a SHIPPING order without "
                    "an order_id."
                )

            if order_state != "SHIPPING":
                raise RuntimeError(
                    f"Mirakl returned order {order_id} with "
                    f"unexpected state {order_state!r}. "
                    "Expected SHIPPING."
                )

            if order_id in seen_order_ids:
                raise RuntimeError(
                    f"Mirakl returned duplicate order "
                    f"{order_id} while paging the queue."
                )

            order_lines = raw_order.get("order_lines", [])

            if not isinstance(order_lines, list) or not order_lines:
                raise RuntimeError(
                    f"Mirakl order {order_id} contains no "
                    "order lines."
                )

            # The existing B&Q processor is designed around one
            # active marketplace product per order. Do not silently
            # guess which line to process if Mirakl returns more.
            if len(order_lines) != 1:
                raise RuntimeError(
                    f"Mirakl order {order_id} contains "
                    f"{len(order_lines)} order lines. "
                    "The current B&Q processor expects exactly 1."
                )

            line = order_lines[0]
            sku = str(
                line.get("offer_sku") or ""
            ).strip()

            if not sku:
                raise RuntimeError(
                    f"Mirakl order {order_id} has no offer SKU."
                )

            customer = raw_order.get("customer") or {}
            shipping = customer.get("shipping_address") or {}

            postcode = str(
                shipping.get("zip_code") or ""
            ).strip()

            if not postcode:
                raise RuntimeError(
                    f"Mirakl order {order_id} has no shipping "
                    "postcode."
                )

            total_price = raw_order.get("total_price")

            try:
                price = float(total_price)
            except (TypeError, ValueError) as exc:
                raise RuntimeError(
                    f"Mirakl order {order_id} has an invalid "
                    f"total price: {total_price!r}."
                ) from exc

            marketplace_orders.append(
                MarketplaceOrder(
                    marketplace="DIY",
                    order_id=order_id,
                    sku=sku,
                    price=price,
                    postcode=postcode,
                    phone_1=str(
                        shipping.get("phone") or ""
                    ).strip(),
                    phone_2=str(
                        shipping.get("phone_secondary") or ""
                    ).strip(),
                )
            )

            seen_order_ids.add(order_id)

        print(
            f"  Page offset {offset}: "
            f"{len(raw_orders)} order(s)"
        )

        # A short page is the final page. An empty page also ends
        # the loop safely.
        if len(raw_orders) < max_results:
            break

        offset += max_results

    print()
    print(
        f"✓ Mirakl returned {len(marketplace_orders)} "
        "B&Q order(s) awaiting shipment."
    )
    print("No Mirakl write has been made.")
    print("=" * 70)

    return marketplace_orders


# ============================================================
# MIRAKL ORDER LOOKUP
# ============================================================

def read_mirakl_order(order_id):
    """
    Read one B&Q/Kingfisher Mirakl order.

    READ ONLY.
    """

    url = (
        f"{MIRAKL_BASE_URL.rstrip('/')}"
        f"/api/orders"
    )

    headers = {
        "Authorization": MIRAKL_API_KEY,
        "Accept": "application/json",
    }

    params = {
        "order_ids": order_id
    }

    response = requests.get(
        url,
        headers=headers,
        params=params,
        timeout=30,
    )

    if response.status_code != 200:

        raise RuntimeError(
            f"Mirakl order lookup failed. "
            f"HTTP {response.status_code}: "
            f"{response.text}"
        )

    data = response.json()

    orders = data.get(
        "orders",
        []
    )

    if len(orders) != 1:

        raise RuntimeError(
            f"Expected exactly 1 Mirakl order, "
            f"found {len(orders)}."
        )

    order = orders[0]

    
    customer = order.get(
        "customer",
        {}
    )

    shipping = customer.get(
        "shipping_address",
        {}
    )

    order_lines = order.get(
        "order_lines",
        []
    )

    
    if not order_lines:

        raise RuntimeError(
            "Mirakl order contains no order lines."
        )

    # Current prototype assumes one active product line.
    line = order_lines[0]

    return {
        "order_id": order.get(
            "order_id"
        ),
        "sku": line.get(
            "offer_sku"
        ),
        "price": float(
            order.get(
                "total_price",
                0
            )
        ),
        "postcode": shipping.get(
            "zip_code"
        ),
        "phone_1": shipping.get(
            "phone"
        ),
        "phone_2": shipping.get(
            "phone_secondary"
        ),
        "customer_name": (
            f"{shipping.get('firstname', '')} "
            f"{shipping.get('lastname', '')}"
        ).strip(),
        "description": line.get(
            "description"
        ),
        "order_state": order.get(
            "order_state"
        ),
        "shipping_carrier_code": order.get(
            "shipping_carrier_code"
        ),
        "shipping_tracking": order.get(
            "shipping_tracking"
        ),
        "shipping_deadline": order.get(
            "shipping_deadline"
        ),
        "delivery_date_earliest": (
                order.get("delivery_date", {})
                .get("earliest")
        ),
        "delivery_date_latest": (
            order.get("delivery_date", {})
            .get("latest")
        ),
    }


# ============================================================
# MIRAKL THREADS
# ============================================================

def clean_mirakl_message_body(body):
    """
    Convert Mirakl HTML message content into readable
    terminal text.
    """

    if not body:
        return ""

    text = body

    text = re.sub(
        r"<br\s*/?>",
        "\n",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"<[^>]+>",
        "",
        text
    )

    text = html.unescape(
        text
    )

    return text.strip()


def read_mirakl_threads(order_id):
    """
    Read existing Mirakl threads/messages for a B&Q order.

    READ ONLY.
    """

    url = (
        f"{MIRAKL_BASE_URL.rstrip('/')}"
        f"/api/inbox/threads"
    )

    headers = {
        "Authorization": MIRAKL_API_KEY,
        "Accept": "application/json",
    }

    params = {
        "entity_type": "MMP_ORDER",
        "entity_id": order_id,
        "with_messages": "true",
    }

    response = requests.get(
        url,
        headers=headers,
        params=params,
        timeout=30,
    )

    if response.status_code != 200:

        raise RuntimeError(
            f"Mirakl thread lookup failed. "
            f"HTTP {response.status_code}: "
            f"{response.text}"
        )

    data = response.json()

    threads = data.get(
        "data",
        []
    )

    all_messages = []

    for thread in threads:

        topic = thread.get(
            "topic",
            {}
        )

        for message in thread.get(
            "messages",
            []
        ):

            sender = message.get(
                "from",
                {}
            )

            all_messages.append(
                {
                    "thread_id": thread.get(
                        "id"
                    ),
                    "topic_type": topic.get(
                        "type"
                    ),
                    "topic_code": topic.get(
                        "value"
                    ),
                    "message_id": message.get(
                        "id"
                    ),
                    "date": message.get(
                        "date_created"
                    ),
                    "sender_type": sender.get(
                        "type"
                    ),
                    "sender_name": sender.get(
                        "display_name"
                    ),
                    "body": clean_mirakl_message_body(
                        message.get(
                            "body"
                        )
                    ),
                }
            )

    all_messages.sort(
        key=lambda message: (
            message["date"] or ""
        )
    )

    return {
        "threads": threads,
        "messages": all_messages,
    }


# ============================================================
# MIRAKL DRY-RUN PACKAGE
# ============================================================

def build_mirakl_dry_run_payload(
    order_id,
    customer_name,
    message_body,
    invoice_path,
    existing_threads,
):
    """
    Build our LOCAL representation of the Mirakl delivery
    message.

    IMPORTANT:
    No network write occurs here.
    """

    thread_id = None

    if existing_threads:
        thread_id = existing_threads[0].get(
            "id"
        )

    return {
        "mode": "DRY_RUN",
        "order_id": order_id,
        "customer": customer_name,
        "topic": {
            "type": "REASON_CODE",
            "value": DIY_DELIVERY_TOPIC_CODE,
            "label": (
                "Information about delivery "
                "(incl. tracking)"
            ),
        },
        "existing_thread_id": thread_id,
        "message_body": message_body,
        "attachment": {
            "filename": os.path.basename(
                invoice_path
            ),
            "path": invoice_path,
            "exists": os.path.isfile(
                invoice_path
            ),
        },
    }

def format_mirakl_message_html(message_body):
    """
    Convert our plain-text customer message into safe HTML
    for Mirakl while preserving line breaks and paragraphs.
    """

    if not message_body:
        return ""

    # Escape customer/order text so characters such as
    # <, > and & cannot accidentally become HTML.
    safe_body = html.escape(
        message_body
    )

    # Mirakl renders the message body as HTML, so normal
    # newline characters need explicit HTML line breaks.
    safe_body = safe_body.replace(
        "\n",
        "<br />"
    )

    return safe_body

def build_mirakl_shipping_dry_run(
    mirakl_order,
):
    """
    Build a local representation of the proposed
    Mirakl shipping action.

    NO network write occurs here.
    """

    order_id = mirakl_order["order_id"]

    return {
        "mode": "DRY_RUN",
        "order_id": order_id,
        "current_state": mirakl_order["order_state"],
        "target_state": "SHIPPED",

        "tracking_number": order_id,

        "carrier_code": (
            DIY_DIRECT_DELIVERY_CARRIER_CODE
        ),

        "carrier_label": (
            DIY_DIRECT_DELIVERY_CARRIER_LABEL
        ),

        "endpoint": (
            f"/api/orders/"
            f"{order_id}/ship"
        ),
    }

def ship_mirakl_order(
    order_id,
    carrier_code,
    carrier_name,
    tracking_number,
):
    """
    Set tracking/carrier details, then validate shipment.

    IMPORTANT:
    Two PUT requests are made:
    1. Update tracking/carrier
    2. Mark order as shipped

    There is deliberately NO automatic retry.
    """

    headers = {
        "Authorization": MIRAKL_API_KEY,
        "Accept": "application/json",
        "Content-Type": "application/json",
    }

    # --------------------------------------------------------
    # STEP 1 - UPDATE TRACKING / CARRIER
    # --------------------------------------------------------

    tracking_url = (
        f"{MIRAKL_BASE_URL.rstrip('/')}"
        f"/api/orders/{order_id}/tracking"
    )

    tracking_payload = {
        "carrier_code": carrier_code,
        "carrier_name": carrier_name,
        "tracking_number": tracking_number,
    }

    tracking_response = requests.put(
        tracking_url,
        headers=headers,
        json=tracking_payload,
        timeout=30,
    )

    if not tracking_response.ok:

        raise RuntimeError(
            f"Mirakl tracking update failed. "
            f"HTTP {tracking_response.status_code}: "
            f"{tracking_response.text}"
        )

    # --------------------------------------------------------
    # STEP 2 - MARK ORDER AS SHIPPED
    # --------------------------------------------------------

    ship_url = (
        f"{MIRAKL_BASE_URL.rstrip('/')}"
        f"/api/orders/{order_id}/ship"
    )

    ship_response = requests.put(
        ship_url,
        headers=headers,
        timeout=30,
    )

    if not ship_response.ok:

        raise RuntimeError(
            f"Mirakl ship request failed. "
            f"HTTP {ship_response.status_code}: "
            f"{ship_response.text}"
        )

    return {
        "success": True,
        "tracking_status_code": (
            tracking_response.status_code
        ),
        "ship_status_code": (
            ship_response.status_code
        ),
        "tracking_response_text": (
            tracking_response.text
        ),
        "ship_response_text": (
            ship_response.text
        ),
    }

def send_mirakl_delivery_message(
    order_id,
    message_body,
    invoice_path,
):
    """
    Send one delivery-information thread to the customer
    through Mirakl.

    IMPORTANT:
    This performs ONE POST request only.
    There is deliberately NO automatic retry.
    """

    url = (
        f"{MIRAKL_BASE_URL.rstrip('/')}"
        f"/api/orders/{order_id}/threads"
    )

    headers = {
        "Authorization": MIRAKL_API_KEY,
        "Accept": "application/json",
    }

    formatted_message_body = format_mirakl_message_html(
    message_body
)

    thread_input = {
        "body": formatted_message_body,
        "topic": {
            "type": "REASON_CODE",
            "value": DIY_DELIVERY_TOPIC_CODE,
        },
        "to": [
            "CUSTOMER",
        ],
    }

    if not os.path.isfile(invoice_path):
        raise RuntimeError(
            f"Invoice attachment does not exist: "
            f"{invoice_path}"
        )

    with open(invoice_path, "rb") as invoice_file:

        files = {
            "files": (
                os.path.basename(invoice_path),
                invoice_file,
                "application/pdf",
            ),

            "thread_input": (
                None,
                json.dumps(thread_input),
                "application/json",
            ),
        }

        response = requests.post(
            url,
            headers=headers,
            files=files,
            timeout=30,
        )

    if not response.ok:

        raise RuntimeError(
            f"Mirakl message send failed. "
            f"HTTP {response.status_code}: "
            f"{response.text}"
        )

    try:
        response_data = response.json()

    except ValueError:
        response_data = {
            "raw_response": response.text
        }

    return {
        "success": True,
        "status_code": response.status_code,
        "response": response_data,
    }

class DIYMarketplaceAdapter(MarketplaceAdapter):
    """
    DIY / B&Q marketplace adapter.
    """

    def __init__(self, orders_file):
        self.orders_file = orders_file

    @property
    def name(self) -> str:
        return "DIY"

    def read_orders(self):
        """
        Return all currently exported DIY orders.
        """
        return read_bq_orders(
            self.orders_file
        )

    def read_first_order(self) -> MarketplaceOrder:
        """
        Compatibility method.
        """
        return read_first_bq_order(
            self.orders_file
        )