import os
import re
from datetime import datetime, timedelta

from dotenv import load_dotenv
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

# Change this depending on where you are running the program:
#
# "work"   = work network
# "remote" = home / remote connection
#
MODE = "remote"
ORDERS_FILE = "orders.xlsx"

WORK_URL = "https://gar.rpdns.co.uk/companyGAR/rp2.cgi"
REMOTE_URL = "https://garremote.rpdns.co.uk/companyGAR/rp2.cgi"
INVOICE_FOLDER = r"C:\Users\Noel\Documents\OK TO DELETE"

# Temporary test order.
# We will replace this with the B&Q Excel import next.
TEST_ORDER = None

INVOICE_FOLDER = r"C:\Users\Noel\Documents\OK TO DELETE"



# ============================================================
# LOAD REMOTE CREDENTIALS
# ============================================================

load_dotenv()

RP2_REMOTE_USERNAME = os.getenv("RP2_REMOTE_USERNAME")
RP2_REMOTE_PASSWORD = os.getenv("RP2_REMOTE_PASSWORD")


if MODE == "remote":
    RP2_URL = REMOTE_URL

    if not RP2_REMOTE_USERNAME or not RP2_REMOTE_PASSWORD:
        raise RuntimeError(
            "Remote RPii username/password missing from .env file."
        )

elif MODE == "work":
    RP2_URL = WORK_URL

else:
    raise ValueError(
        'MODE must be either "work" or "remote".'
    )


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def read_payment_summary(sales_frame):
    """
    Reads the RPii Payment Summary.

    A sale is considered payment-processed if the Payment Summary
    contains at least one payment row beneath the heading.

    The payment amount itself is NOT used to determine this.
    PREPAL, for example, may legitimately show 0.00.
    """

    payment_table = sales_frame.locator("#psum")

    if payment_table.count() == 0:
        return {
            "processed": False,
            "details": None,
        }

    rows = payment_table.locator("tr")

    payment_details = []

    for i in range(rows.count()):
        row = rows.nth(i)
        text = " ".join(
            row.inner_text().split()
        )

        # Ignore the heading
        if not text:
            continue

        if text.lower() == "payment summary":
            continue

        payment_details.append(text)

    if payment_details:
        return {
            "processed": True,
            "details": payment_details,
        }

    return {
        "processed": False,
        "details": None,
    }

def extract_web_ref(text):
    """
    Extracts:

        Web Ref: 1068814250-A

    from the larger RPii sale information panel.
    """

    match = re.search(
        r"Web Ref:\s*([A-Za-z0-9\-]+)",
        text
    )

    return match.group(1) if match else None


def extract_customer_name(text):
    """
    RPii customer panel example:

        DIY0205 • DIY • ...

    For B&Q orders this is normally the generic
    RPii account such as DIY, rather than the actual
    delivery customer's name.
    """

    parts = [
        part.strip()
        for part in text.split("•")
    ]

    if len(parts) >= 2:
        return parts[1]

    return None

def get_customer_surname(full_name):
    """
    Extract a safe surname for use in an invoice filename.

    Example:
        "mark booth" -> "Booth"
        "Jackie Jeeves" -> "Jeeves"
    """

    if not full_name:
        return "Unknown"

    parts = full_name.strip().split()

    if not parts:
        return "Unknown"

    surname = parts[-1]

    # Remove characters Windows does not allow in filenames
    surname = re.sub(
        r'[<>:"/\\|?*]',
        '',
        surname
    )

    return surname.title()

def extract_delivery_customer(text):
    """
    Example RPii text:

    JEEV9737 • Jackie Jeeves •
    36, Smith Square, Doncaster, DN4 0SR
    T: 07761022352 • e: ...
    """

    cleaned = " ".join(text.split())

    parts = [
        part.strip()
        for part in cleaned.split("•")
    ]

    name = None
    address = None
    telephone = None

    if len(parts) >= 2:
        name = parts[1]

    if len(parts) >= 3:
        address = parts[2]

        # Remove telephone/email text if it has been folded into this part
        address = re.split(
            r"\bT:\s*",
            address
        )[0].strip()

    tel_match = re.search(
        r"T:\s*([0-9 ]+)",
        cleaned
    )

    if tel_match:
        telephone = tel_match.group(1).strip()

    return {
        "name": name,
        "address": address,
        "telephone": telephone,
    }

def extract_uk_postcode(address):
    """
    Extract a UK postcode from the end of an address.

    Examples:
        DN4 0SR
        BB18 6PZ
        SW1A 1AA
    """

    if not address:
        return None

    match = re.search(
        r"\b("
        r"[A-Z]{1,2}\d[A-Z\d]?"
        r"\s*"
        r"\d[A-Z]{2}"
        r")\b",
        address.upper()
    )

    if not match:
        return None

    postcode = match.group(1).replace(" ", "")

    # Standard UK postcode spacing:
    # everything except last 3 characters + space + last 3
    return postcode[:-3] + " " + postcode[-3:]

def normalise_phone(number):
    if not number:
        return ""

    number = re.sub(r"\D", "", str(number))

    if number.startswith("44"):
        number = "0" + number[2:]

    return number

def parse_delivery_block(text):
    """
    Example SGK:

        DELIVER SK
        BN:09
        070926:AM
        SK

    Example EDUK:

        DELIVER ED
        BN:09
        120926:PM
        ED

    Business rule:

        SGK:
            use the RPii date exactly.

        ED:
            actual customer delivery date is
            ONE DAY BEFORE the RPii date.
    """

    cleaned = " ".join(text.split())

    courier = None

    if "DELIVER SK" in cleaned:
        courier = "SK"

    elif "DELIVER ED" in cleaned:
        courier = "ED"

    date_match = re.search(
        r"(\d{6}):(AM|PM)",
        cleaned
    )

    rp2_date = None
    customer_date = None
    slot = None

    if date_match:
        date_text = date_match.group(1)
        slot = date_match.group(2)

        rp2_date = datetime.strptime(
            date_text,
            "%d%m%y"
        ).date()

        if courier == "ED":
            customer_date = (
                rp2_date - timedelta(days=1)
            )

        else:
            customer_date = rp2_date

    return {
        "courier": courier,
        "rp2_date": rp2_date,
        "customer_date": customer_date,
        "slot": slot,
        "raw": cleaned,
    }


def read_active_product(sales_frame):
    """
    RPii presents sale product information in pairs:

        Cell 0 = SKU / stock information
        Cell 1 = description
        Cell 2 = next SKU
        Cell 3 = next description

    Example cancelled/import-problem line:

        WEBPART ( - : )

    Example active line:

        IB55732W (IB55 732 W UK) (1 -S L:W2)

    We therefore choose the first product line
    containing a positive quantity.
    """

    product_cells = sales_frame.locator(
        "td.SALES-lines-left-wrap"
    )

    product_cell_count = product_cells.count()

    sku = None
    description = None
    quantity = None

    for i in range(
        0,
        product_cell_count - 1,
        2
    ):
        sku_cell = product_cells.nth(i)
        desc_cell = product_cells.nth(i + 1)

        sku_text = " ".join(
            sku_cell.inner_text().split()
        )

        description_text = " ".join(
            desc_cell.inner_text().split()
        )

        # Active product example:
        #
        # (1 -S L:W2)
        #
        active_match = re.search(
            r"\(\s*(\d+)\s+-",
            sku_text
        )

        if not active_match:
            continue

        candidate_quantity = int(
            active_match.group(1)
        )

        if candidate_quantity <= 0:
            continue

        candidate_sku = sku_text.split()[0]

        sku = candidate_sku
        description = description_text
        quantity = candidate_quantity

        break

    return {
        "sku": sku,
        "description": description,
        "quantity": quantity,
    }

def read_first_bq_order(filename):
    workbook = load_workbook(
        filename,
        data_only=True
    )

    sheet = workbook.active

    headers = {
        cell.value: cell.column
        for cell in sheet[1]
        if cell.value
    }

    required_columns = [
        "Order number",
        "Offer SKU",
        "Amount",
        "Billing address zip",
        "Shipping address phone",
        "Shipping address phone 2",
    ]

    for column in required_columns:
        if column not in headers:
            raise RuntimeError(
                f"Missing required Excel column: {column}"
            )

    # First data row = oldest/top order
    row = 2

    order_id = sheet.cell(
        row=row,
        column=headers["Order number"]
    ).value

    sku = sheet.cell(
        row=row,
        column=headers["Offer SKU"]
    ).value

    price = sheet.cell(
        row=row,
        column=headers["Amount"]
    ).value

    postcode = sheet.cell(
        row=row,
        column=headers["Billing address zip"]
    ).value

    phone_1 = sheet.cell(
        row=row,
        column=headers["Shipping address phone"]
    ).value

    phone_2 = sheet.cell(
    row=row,
        column=headers["Shipping address phone 2"]
    ).value

    return {
        "order_id": str(order_id).strip(),
        "sku": str(sku).strip(),
        "price": float(price),
        "postcode": str(postcode).strip(),
        "phone_1": str(phone_1).strip() if phone_1 else "",
        "phone_2": str(phone_2).strip() if phone_2 else "",
    }


excel_order = read_first_bq_order(
        ORDERS_FILE
    )

print("\nEXCEL DEBUG:")
print(excel_order)

TEST_ORDER = excel_order["order_id"]

print()
print("=" * 60)
print("B&Q EXCEL ORDER")
print("=" * 60)
print(
    f"Order ID:      {excel_order['order_id']}"
)
print(
    f"SKU:           {excel_order['sku']}"
)
print(
    f"Price:         £{excel_order['price']:.2f}"
)
print("=" * 60)

input(
    "\nCheck the Excel order above. "
    "Press ENTER to continue..."
)


# ============================================================
# START PLAYWRIGHT
# ============================================================

with sync_playwright() as p:

    browser = p.chromium.launch(
        headless=False
    )

    # --------------------------------------------------------
    # Create browser context
    # --------------------------------------------------------

    if MODE == "remote":

        context = browser.new_context(
            http_credentials={
                "username": RP2_REMOTE_USERNAME,
                "password": RP2_REMOTE_PASSWORD,
            }
        )

    else:

        context = browser.new_context()

    page = context.new_page()

    # --------------------------------------------------------
    # Open RPii
    # --------------------------------------------------------

    print()
    print("=" * 60)
    print("EDUK RPii ORDER READER")
    print("=" * 60)
    print(f"Mode: {MODE.upper()}")
    print(f"Opening: {RP2_URL}")
    print()

        # --------------------------------------------------------
    # Open RPii and wait for the actual EPoS module
    # --------------------------------------------------------

    print("Opening RPii...")

    max_attempts = 3
    rp2_ready = False

    for attempt in range(1, max_attempts + 1):

        print(
            f"\nRPii load attempt "
            f"{attempt}/{max_attempts}..."
        )

        try:
            page.goto(
                RP2_URL,
                wait_until="domcontentloaded",
                timeout=30000
            )

            print(
                "Outer RPii page loaded. "
                "Waiting for Sales module..."
            )

            # RPii remote can be slow to initialise
            page.wait_for_selector(
                "#SALESFrame",
                state="attached",
                timeout=30000
            )

            sales_frame = page.frame_locator(
                "#SALESFrame"
            )

            sales_enquiry = sales_frame.locator(
                "#orderno"
            )

            # Give remote RPii plenty of time
            sales_enquiry.wait_for(
                state="visible",
                timeout=60000
            )

            rp2_ready = True

            print(
                "✓ RPii Sales module is ready."
            )

            break

        except Exception as error:

            print(
                "RPii did not finish loading."
            )

            print(
                f"Reason: {error}"
            )

            if attempt < max_attempts:

                print(
                    "Waiting 5 seconds before retry..."
                )

                page.wait_for_timeout(5000)

            else:

                print()
                print("=" * 60)
                print(
                    "RPii REMOTE CONNECTION FAILED"
                )
                print("=" * 60)
                print(
                    "The outer RPii screen loaded, "
                    "but the Sales module did not."
                )
                print()
                print(
                    "No order processing has taken place."
                )

    if not rp2_ready:

        input(
            "\nPress ENTER to close..."
        )

        browser.close()

        raise SystemExit

    input(
        "\nRPii is ready. "
        "Press ENTER to open the test order..."
    )

    # --------------------------------------------------------
    # Search for order
    # --------------------------------------------------------

    print(
        f"\nOpening order {TEST_ORDER}..."
    )

    sales_enquiry.fill(TEST_ORDER)
    sales_enquiry.press("Enter")

    # Wait until the sale total has appeared.
    # This is better than relying only on a fixed delay.
    total_sale_field = sales_frame.locator(
        "#qh_ordval"
    )

    total_sale_field.wait_for(
        state="visible",
        timeout=30000
    )

    # Small additional pause because RPii updates
    # several parts of the sale asynchronously.
    page.wait_for_timeout(1000)

    # ========================================================
    # CUSTOMER PANEL DIAGNOSTIC
    # ========================================================

    print()
    print("CUSTOMER PANEL DIAGNOSTIC")
    print("=" * 70)

    customer_panels = sales_frame.locator(
        "td.SALES-panel-s"
    )

    print(
        f"SALES-panel-s elements found: "
        f"{customer_panels.count()}"
    )

    for i in range(customer_panels.count()):

        panel = customer_panels.nth(i)

        print()
        print(f"PANEL {i}")
        print("TEXT:")
        print(repr(panel.inner_text()))
        print("-" * 50)

    print("=" * 70)

    # ========================================================
    # READ WEB REFERENCE
    # ========================================================

    web_panel = sales_frame.locator(
        "td.SALES-panel",
        has_text="Web Ref:"
    )

    web_panel_text = web_panel.inner_text()

    web_ref = extract_web_ref(
        web_panel_text
    )

    # ========================================================
    # SAFETY CHECK 1 - ORDER NUMBER
    # ========================================================

    print()
    print("=" * 60)
    print("SAFETY CHECK 1 - ORDER NUMBER")
    print("=" * 60)

    print(
        f"Excel Order:   {excel_order['order_id']}"
    )

    print(
        f"RPii Web Ref:  {web_ref}"
    )

    if web_ref == excel_order["order_id"]:

        print()
        print("✓ ORDER NUMBER MATCH")

    else:

        print()
        print("✗ ORDER NUMBER MISMATCH")
        print()
        print(
            "STOPPING - no further processing will take place."
        )

        input(
            "\nPress ENTER to close..."
        )

        browser.close()

        raise SystemExit

    

    # ========================================================
    # READ RPii CUSTOMER ACCOUNT
    # ========================================================

    ccustomer_panels = sales_frame.locator(
    "td.SALES-panel-s"
)

    customer_panel = None

    for i in range(customer_panels.count()):
        panel = customer_panels.nth(i)
        text = panel.inner_text()

        if "T:" in text:
            customer_panel = panel
            break

    if customer_panel is None:
        raise RuntimeError(
            "Could not find RPii delivery customer panel."
        )

    customer_text = customer_panel.inner_text()

    customer_name = extract_customer_name(
        customer_text
    )

    # ========================================================
    # READ DELIVERY CUSTOMER DETAILS
    # ========================================================

    delivery_customer = extract_delivery_customer(
        customer_text
    )

    delivery_postcode = extract_uk_postcode(
       delivery_customer["address"]
    )

    # ========================================================
    # SAFETY CHECK 4 - POSTCODE
    # ========================================================

    print()
    print("=" * 60)
    print("SAFETY CHECK 4 - POSTCODE")
    print("=" * 60)

    excel_postcode = (
        excel_order["postcode"]
        .replace(" ", "")
        .upper()
    )

    rp2_postcode = (
        delivery_postcode
        .replace(" ", "")
        .upper()
        if delivery_postcode
        else None
    )

    print(
        f"Excel Postcode: {excel_order['postcode']}"
    )

    print(
        f"RPii Postcode:  {delivery_postcode}"
    )

    if rp2_postcode == excel_postcode:

        print()
        print("✓ POSTCODE MATCH")

    else:

        print()
        print("✗ POSTCODE MISMATCH")
        print()
        print(
            "STOPPING - no further processing will take place."
        )

        input(
            "\nPress ENTER to close..."
        )

        browser.close()

        raise SystemExit
    
    # ========================================================
    # READ ACTIVE PRODUCT
    # ========================================================

    product = read_active_product(
        sales_frame
    )

    sku = product["sku"]
    description = product["description"]
    quantity = product["quantity"]

    # ========================================================
    # SAFETY CHECK 2 - SKU
    # ========================================================
    
    print()
    print("=" * 60)
    print("SAFETY CHECK 2 - SKU")
    print("=" * 60)

    excel_sku = str(
        excel_order["sku"]
    ).strip().upper()

    rp2_sku = str(
        sku
    ).strip().upper()

    print(
        f"Excel SKU:     {excel_sku}"
    )

    print(
        f"RPii SKU:      {rp2_sku}"
    )

    if excel_sku == rp2_sku:

        print()
        print("✓ SKU MATCH")

    else:

        print()
        print("✗ SKU MISMATCH")
        print()
        print(
            "STOPPING - no further processing will take place."
        )

        input(
            "\nPress ENTER to close..."
        )

        browser.close()

        raise SystemExit

    # ========================================================
    # READ FINANCIALS
    # ========================================================

    total_sale = sales_frame.locator(
        "#qh_ordval"
    ).input_value()

    total_paid = sales_frame.locator(
        "#qh_totpaid"
    ).input_value()

    balance = sales_frame.locator(
        "#qh_balance"
    ).input_value()

    # ========================================================
    # READ PAYMENT SUMMARY
    # ========================================================

    payment_summary = read_payment_summary(
        sales_frame
    )

    # ========================================================
    # SAFETY CHECK 3 - ORDER VALUE
    # ========================================================

    print()
    print("=" * 60)
    print("SAFETY CHECK 3 - ORDER VALUE")
    print("=" * 60)

    excel_amount = round(
        float(excel_order["price"]),
        2
    )

    rp2_amount = round(
        float(total_sale),
        2
    )

    print(
        f"Excel Amount:  £{excel_amount:.2f}"
    )

    print(
        f"RPii Total:    £{rp2_amount:.2f}"
    )

    if excel_amount == rp2_amount:

        print()
        print("✓ ORDER VALUE MATCH")

    else:

        print()
        print("✗ ORDER VALUE MISMATCH")
        print()
        print(
            "STOPPING - no further processing will take place."
        )

        input(
            "\nPress ENTER to close..."
        )

        browser.close()

        raise SystemExit

    # ========================================================
    # READ COURIER / DELIVERY INFORMATION
    # ========================================================

    delivery_cells = sales_frame.locator(
        "td.tcb"
    )

    delivery_text = None

    for i in range(
        delivery_cells.count()
    ):
        text = delivery_cells.nth(
            i
        ).inner_text()

        if "DELIVER" in text:
            delivery_text = text
            break

    if delivery_text:

        delivery = parse_delivery_block(
            delivery_text
        )

    else:

        delivery = {
            "courier": None,
            "rp2_date": None,
            "customer_date": None,
            "slot": None,
            "raw": None,
        }

    # ========================================================
    # REPORT
    # ========================================================

    print()
    print()
    print("=" * 60)
    print("RPii ORDER CHECK")
    print("=" * 60)

    print(
        f"Web Ref:       {web_ref}"
    )

    print(
        f"RPii Account:  {customer_name}"
    )

    print(
        f"SKU:           {sku}"
    )

    print(
        f"Description:   {description}"
    )

    if quantity is not None:
        print(
            f"Quantity:      {quantity}"
        )

    print()

    print(
        f"Total Sale:    £{total_sale}"
    )

    print(
        f"Total Paid:    £{total_paid}"
    )

    print(
        f"Balance:       £{balance}"
    )

    print()

    courier = delivery["courier"]

    # --------------------------------------------------------
    # Payment reporting
    # --------------------------------------------------------
    
    print()
    print("PAYMENT SUMMARY")
    print("-" * 60)

    if payment_summary["processed"]:

        print("Status:        PAYMENT PRESENT")

        for detail in payment_summary["details"]:
            print(
                f"Details:       {detail}"
            )

    else:

        print("Status:        PAYMENT REQUIRED")
        print("Details:       None")

    # --------------------------------------------------------
    # Courier reporting
    # --------------------------------------------------------

    if courier == "SK":

        print(
            "Courier:       SGK Distribution"
        )

    elif courier == "ED":

        print(
            "Courier:       Electrical Discount UK"
        )

    else:

        print(
            f"Courier:       UNKNOWN ({courier})"
        )

    # ========================================================
    # DETERMINE ORDER ROUTE
    # ========================================================

    if payment_summary["processed"]:
        order_route = "READY FOR INVOICE"
    else:
        order_route = "WAITING FOR PAYMENT"

    # ========================================================
    # NORMAL ROUTE - OPEN A4 INVOICE
    # ========================================================

    if order_route == "READY FOR INVOICE":

        print()
        print("NORMAL ROUTE")
        print("=" * 60)
        print("Order is ready for invoice.")

        # ----------------------------------------------------
        # Find RPii Print button
        # ----------------------------------------------------

        print_button = sales_frame.locator(
            'td.tcbw[onclick^="doPrint("]'
        )

        if print_button.count() != 1:

            print(
                f"✗ Expected 1 Print button, "
                f"found {print_button.count()}."
            )

            input(
                "\nPress ENTER to close..."
            )

            browser.close()
            raise SystemExit

        # ----------------------------------------------------
        # Open Print options
        # ----------------------------------------------------

        print("Opening RPii Print panel...")

        print_button.evaluate(
            "element => element.click()"
        )

        page.wait_for_timeout(1000)

        print("✓ Print panel opened.")

        # ----------------------------------------------------
        # Find A4 Invoice
        # ----------------------------------------------------

        a4_invoice = sales_frame.locator(
            'tr:has-text("A4 Invoice") '
            'img[onclick^="doInvoiceP("]'
        )

        if a4_invoice.count() != 1:

            print(
                f"✗ Expected 1 A4 Invoice option, "
                f"found {a4_invoice.count()}."
            )

            input(
                "\nPress ENTER to close..."
            )

            browser.close()
            raise SystemExit

        # Show us the RPii action for diagnostic purposes
        invoice_action = a4_invoice.get_attribute(
            "onclick"
        )

        print(
            f"A4 Invoice action: {invoice_action}"
        )

        # ----------------------------------------------------
        # Open A4 Invoice
        # ----------------------------------------------------

        print("Opening A4 Invoice...")

        a4_invoice.evaluate(
            "element => element.click()"
        )

        page.wait_for_timeout(1500)

        # ========================================================
        # SAVE GENERATED INVOICE PDF
        # ========================================================

        print()
        print("SAVING INVOICE")
        print("=" * 60)

        # Find RPii's PRINTFrame
        print_frame = None

        for frame in page.frames:
            if frame.name == "PRINTFrame":
                print_frame = frame
                break

        if print_frame is None:
            print("✗ PRINTFrame could not be found.")

            input(
                "\nPress ENTER to close..."
            )

            browser.close()
            raise SystemExit

        invoice_url = print_frame.url

        print(
            f"Invoice URL: {invoice_url}"
        )

        # Safety check - make sure this really is a PDF
        if not invoice_url.lower().endswith(".pdf"):
            print(
                "✗ PRINTFrame does not contain a PDF."
            )

            input(
                "\nPress ENTER to close..."
            )

            browser.close()
            raise SystemExit

        # Make sure invoice folder exists
        os.makedirs(
            INVOICE_FOLDER,
            exist_ok=True
        )

        # Create invoice filename
        customer_surname = get_customer_surname(
            delivery_customer["name"]
        )

        invoice_filename = (
            f"{excel_order['order_id']}-"
            f"{customer_surname}-invoice.pdf"
        )

        invoice_path = os.path.join(
            INVOICE_FOLDER,
            invoice_filename
        )

        print(
            f"Saving as: {invoice_filename}"
        )

        # Download using the authenticated
        # Playwright browser context
        response = context.request.get(
            invoice_url
        )

        if not response.ok:

            print(
                f"✗ Invoice download failed. "
                f"HTTP status: {response.status}"
            )

            input(
                "\nPress ENTER to close..."
            )

            browser.close()
            raise SystemExit

        # Save PDF
        with open(
            invoice_path,
            "wb"
        ) as file:

            file.write(
                response.body()
            )

        print()
        print("✓ INVOICE SAVED SUCCESSFULLY")
        print()
        print(
            f"Saved to: {invoice_path}"
        )

        input(
            "\nCheck the saved PDF file. "
            "Press ENTER to close..."
        )

        browser.close()
        raise SystemExit

    # --------------------------------------------------------
    # RPii scheduled date
    # --------------------------------------------------------

    if delivery["rp2_date"]:

        print(
            "RPii Date:      "
            + delivery["rp2_date"].strftime(
                "%d/%m/%Y"
            )
        )

    print()
    print("DELIVERY CUSTOMER")
    print("-" * 60)

    print(
        f"Name:          {delivery_customer['name']}"
    )

    print(
        f"Address:       {delivery_customer['address']}"
    )

    print(
       f"Postcode:      {delivery_postcode}"
    )

    print(
        f"Telephone:     {delivery_customer['telephone']}"
    )

    print()
    print("ORDER ROUTE")
    print("-" * 60)
    print(f"Route:         {order_route}")

    # ========================================================
    # SAFETY CHECK 5 - TELEPHONE
    # ========================================================

    print()
    print("=" * 60)
    print("SAFETY CHECK 5 - TELEPHONE")
    print("=" * 60)

    excel_phone_1 = normalise_phone(
        excel_order["phone_1"]
    )

    excel_phone_2 = normalise_phone(
        excel_order["phone_2"]
    )

    rp2_phone = normalise_phone(
        delivery_customer["telephone"]
    )

    print(
        f"Excel Phone 1: {excel_order['phone_1']}"
    )

    print(
        f"Excel Phone 2: {excel_order['phone_2']}"
    )

    print(
        f"RPii Phone:    {delivery_customer['telephone']}"
    )

    valid_excel_phones = {
        phone
        for phone in (
            excel_phone_1,
            excel_phone_2,
        )
        if phone
    }

    if rp2_phone in valid_excel_phones:

        print()
        print("✓ TELEPHONE MATCH")

    else:

        print()
        print("✗ TELEPHONE MISMATCH")
        print()
        print(
            "STOPPING - no further processing will take place."
        )

        input(
            "\nPress ENTER to close..."
        )

        browser.close()

        raise SystemExit

    # --------------------------------------------------------
    # Customer-facing date
    # --------------------------------------------------------

    if delivery["customer_date"]:

        if courier == "SK":

            print(
                "Collection:     "
                + delivery[
                    "customer_date"
                ].strftime(
                    "%d/%m/%Y"
                )
            )

        elif courier == "ED":

            print(
                "Delivery Date:  "
                + delivery[
                    "customer_date"
                ].strftime(
                    "%d/%m/%Y"
                )
            )

            print(
                "                "
                "(RPii date minus 1 day)"
            )

    # --------------------------------------------------------
    # AM / PM
    # --------------------------------------------------------

    if delivery["slot"]:

        print(
            f"Slot:          "
            f"{delivery['slot']}"
        )

    print("=" * 60)

    # ========================================================
    # BASIC SAFETY CHECK
    # ========================================================

    print()
    print("SAFETY CHECK")
    print("-" * 60)

    if web_ref == TEST_ORDER:

        print(
            "✓ Web Ref matches requested order."
        )

    else:

        print(
            "✗ WARNING: Web Ref does NOT match "
            "requested order."
        )

    if sku:

        print(
            f"✓ Active product found: {sku}"
        )

    else:

        print(
            "✗ WARNING: No active product "
            "could be identified."
        )

    if courier in ("SK", "ED"):

        print(
            "✓ Recognised delivery method."
        )

    else:

        print(
            "⚠ Delivery method not recognised."
        )

    print("-" * 60)

    print(
        "\nREAD-ONLY MODE: "
        "No changes have been made to the sale."
    )

    input(
        "\nPress ENTER to close..."
    )

    browser.close() 